"""Shared test fixtures and lightweight fakes.

The fakes here let the pipeline and retriever be tested without downloading a
sentence-transformers model or calling any network service.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest

from app.rag.embeddings import Embedder


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Write a small multi-section ``.docx`` and return its path."""
    from docx import Document

    document = Document()
    document.add_heading("Refund Policy", level=1)
    document.add_paragraph(
        "Customers may request a refund within 30 days of purchase. "
        "Refunds are issued to the original payment method within five business days."
    )
    document.add_heading("Shipping", level=1)
    document.add_paragraph(
        "Standard shipping takes three to five business days. "
        "Express shipping is delivered next day for an extra fee."
    )
    document.add_heading("Support Hours", level=1)
    document.add_paragraph("Support is available Monday to Friday, 9am to 6pm Eastern Time.")

    path = tmp_path / "sample.docx"
    document.save(str(path))
    return path


class FakeEmbedder(Embedder):
    """Deterministic keyword-bag embedder - no model, no torch.

    Each configured keyword owns one dimension; a text's vector counts keyword
    occurrences and is L2-normalized, so texts sharing keywords score highly.
    """

    _VOCAB = ["refund", "shipping", "support", "days", "payment", "express", "hours"]

    def __init__(self) -> None:  # noqa: D107 - see class docstring
        self.model_name = "fake"

    @property
    def dimension(self) -> int:
        return len(self._VOCAB)

    def encode(self, texts: list[str]) -> np.ndarray:
        if not texts:
            return np.zeros((0, self.dimension), dtype=np.float32)
        rows = []
        for text in texts:
            lowered = text.lower()
            vec = np.array([float(lowered.count(word)) for word in self._VOCAB], dtype=np.float32)
            norm = np.linalg.norm(vec)
            rows.append(vec / norm if norm else vec)
        return np.vstack(rows).astype(np.float32)


@pytest.fixture
def fake_embedder() -> FakeEmbedder:
    """A :class:`FakeEmbedder` instance."""
    return FakeEmbedder()
