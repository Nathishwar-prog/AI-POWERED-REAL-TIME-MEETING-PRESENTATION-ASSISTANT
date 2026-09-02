"""Regenerate ``sample_presentation.docx`` used by the demo and manual testing.

Run from the ``backend/`` directory::

    python sample_data/make_sample_docx.py

The generated ``.docx`` is checked in; re-run this only when you want to change
the demo content.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

_OUT = Path(__file__).with_name("sample_presentation.docx")

_SECTIONS: list[tuple[str, list[str]]] = [
    (
        "Product Overview",
        [
            "Northwind Analytics is a self-serve dashboard product for small operations teams.",
            "It connects to spreadsheets and common databases, and refreshes data every 15 minutes.",
            "The product is sold as a monthly subscription with three tiers: Starter, Team, and Business.",
        ],
    ),
    (
        "Pricing and Billing",
        [
            "The Starter tier costs 29 dollars per month and includes 3 dashboards and 1 data source.",
            "The Team tier costs 99 dollars per month and includes 20 dashboards and 5 data sources.",
            "The Business tier costs 299 dollars per month with unlimited dashboards and priority support.",
            "Annual billing receives a 20 percent discount. Invoices are issued on the first of each month.",
        ],
    ),
    (
        "Refund Policy",
        [
            "Customers may request a full refund within 30 days of their first payment.",
            "Refunds after 30 days are prorated for the unused portion of the current billing period.",
            "Refunds are returned to the original payment method within five business days of approval.",
        ],
    ),
    (
        "Security and Data Handling",
        [
            "All data is encrypted in transit with TLS 1.2 or higher and at rest with AES-256.",
            "Northwind Analytics is SOC 2 Type II certified and undergoes an annual third-party audit.",
            "Customer data is stored in the United States and is never sold or shared with advertisers.",
            "Customers can request deletion of their workspace and all associated data at any time.",
        ],
    ),
    (
        "Support and Onboarding",
        [
            "Support is available by email and chat Monday through Friday, 9am to 6pm US Eastern Time.",
            "Business tier customers also receive a dedicated onboarding call and a shared Slack channel.",
            "The median first-response time for support tickets is under two hours during business hours.",
        ],
    ),
    (
        "Roadmap",
        [
            "The next release adds scheduled PDF report delivery and a REST API for exporting metrics.",
            "A native mobile app is planned for later in the year but does not yet have a committed date.",
        ],
    ),
]


def main() -> None:
    document = Document()
    document.add_heading("Northwind Analytics - Customer Briefing", level=0)
    for title, paragraphs in _SECTIONS:
        document.add_heading(title, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
    document.save(str(_OUT))
    print(f"wrote {_OUT}")


if __name__ == "__main__":
    main()
